import io
import zipfile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from app import jobs, providers
from app.limits import BodySizeLimit
from app.main import app
from tests.conftest import FAKE_TRANSCRIPT

client = TestClient(app)

PDF = b'%PDF-1.7\n%fake but correctly signed\n'
WEBM = b'\x1a\x45\xdf\xa3' + bytes(40)
M4A = b'\x00\x00\x00\x20ftypM4A ' + bytes(40)
WAV = b'RIFF\x24\x00\x00\x00WAVEfmt ' + bytes(40)
OGG = b'OggS' + bytes(40)
FLAC = b'fLaC' + bytes(40)
MP3_TAGGED = b'ID3\x04' + bytes(40)
MP3_FRAME = b'\xff\xfb\x90\x64' + bytes(40)


def office_zip(entry: str) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as archive:
        archive.writestr(entry, '<xml/>')
    return buffer.getvalue()


def text_pdf(text: str) -> bytes:
    """A real one-page PDF with the given text, built by hand."""
    stream = f'BT /F1 18 Tf 50 700 Td ({text}) Tj ET'.encode()
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R '
        b'/Resources << /Font << /F1 5 0 R >> >> >>',
        b'<< /Length %d >>\nstream\n' % len(stream) + stream + b'\nendstream',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    out = bytearray(b'%PDF-1.4\n')
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b'%d 0 obj\n' % number + body + b'\nendobj\n'
    xref = len(out)
    out += b'xref\n0 %d\n0000000000 65535 f \n' % (len(objects) + 1)
    out += b''.join(b'%010d 00000 n \n' % offset for offset in offsets)
    out += b'trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n' % (len(objects) + 1, xref)
    return bytes(out)


def text_pptx(title: str, notes: str) -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def post(notes: str = '', files: list[tuple[str, bytes]] | None = None, audio: tuple[str, bytes] | None = None):
    parts = [('slides', (name, content)) for name, content in files or []]
    if audio:
        parts.append(('audio', audio))
    return client.post('/api/sessions', data={'notes': notes}, files=parts or None)


def receipt_of(response):
    assert response.status_code == 202
    return response.json()['receipt']


# The test client runs background jobs before it returns, so the job is already finished here.
def finished_job(response):
    state = client.get(f"/api/jobs/{response.json()['job_id']}").json()
    assert state['status'] in ('done', 'failed')
    return state


def test_health():
    response = client.get('/api/health')
    assert response.status_code == 200
    assert response.json() == {'status': 'ok'}


def test_limits():
    assert client.get('/api/limits').json() == {
        'max_slide_files': 5,
        'max_slide_mb': 25,
        'max_audio_mb': 25,
        'max_notes_chars': 20_000,
        'max_pack_words': 1925,
        'max_pack_chars': 11550,
    }


def test_receipt_for_notes_and_slides():
    pptx = office_zip('ppt/presentation.xml')
    response = post('Asthma: wheeze, cough and chest tightness', [('a.pdf', PDF), ('b.pptx', pptx)])
    assert receipt_of(response) == {
        'notes_words': 6,
        'slides': [
            {'name': 'a.pdf', 'kind': 'pdf', 'size_bytes': len(PDF)},
            {'name': 'b.pptx', 'kind': 'pptx', 'size_bytes': len(pptx)},
        ],
        'audio': None,
    }


def test_audio_formats_from_browsers_are_accepted():
    for name, content, kind in [
        ('rec.webm', WEBM, 'webm'),
        ('rec.m4a', M4A, 'mp4'),
        ('memo.wav', WAV, 'wav'),
        ('rec.ogg', OGG, 'ogg'),
        ('memo.flac', FLAC, 'flac'),
        ('tagged.mp3', MP3_TAGGED, 'mp3'),
        ('bare.mp3', MP3_FRAME, 'mp3'),
    ]:
        assert receipt_of(post(audio=(name, content)))['audio'] == {'name': name, 'kind': kind, 'size_bytes': len(content)}


def test_audio_only_is_enough():
    assert receipt_of(post(audio=('rec.webm', WEBM)))['notes_words'] == 0


def test_renamed_audio_is_rejected():
    assert post(audio=('rec.webm', PDF)).status_code == 415


def test_audio_over_25_mb():
    big = WEBM + bytes(25 * 1024 * 1024)
    assert post(audio=('long.webm', big)).status_code == 413


def test_nothing_to_send():
    assert post('   ').status_code == 422


def test_renamed_file_is_rejected():
    response = post(files=[('slides.pdf', b'MZ\x90\x00 not a pdf')])
    assert response.status_code == 415


def test_word_document_renamed_to_pptx_is_rejected():
    response = post(files=[('slides.pptx', office_zip('word/document.xml'))])
    assert response.status_code == 415


def test_damaged_zip_is_rejected():
    damaged = office_zip('ppt/presentation.xml')[:-10]
    assert post(files=[('slides.pptx', damaged)]).status_code == 415


def test_pdf_signature_after_leading_bytes():
    assert receipt_of(post(files=[('a.pdf', bytes(20) + PDF)]))['slides'][0]['kind'] == 'pdf'


def test_too_many_files():
    files = [(f'{n}.pdf', PDF) for n in range(6)]
    assert post(files=files).status_code == 422


def test_file_over_25_mb():
    big = PDF + b'0' * (25 * 1024 * 1024)
    assert post(files=[('big.pdf', big)]).status_code == 413


def test_path_in_filename_is_stripped():
    assert receipt_of(post(files=[('../../etc/a.pdf', PDF)]))['slides'][0]['name'] == 'a.pdf'


def test_job_reads_slides_and_transcribes(fake_transcribe):
    pdf = text_pdf('Peak flow below 50 percent means severe asthma')
    pptx = text_pptx('Acute asthma', 'Ask about previous ICU admissions')
    response = post('my notes', [('lecture.pdf', pdf), ('round.pptx', pptx)], audio=('round.webm', WEBM))

    state = finished_job(response)
    assert state['status'] == 'done'
    materials = state['materials']
    assert materials['transcript'] == FAKE_TRANSCRIPT
    assert 'Peak flow below 50 percent' in materials['slides'][0]['text']
    assert materials['slides'][0]['pages'] == 1
    assert 'Acute asthma' in materials['slides'][1]['text']
    assert 'Speaker notes: Ask about previous ICU admissions' in materials['slides'][1]['text']
    assert materials['warnings'] == []


def test_audio_reaches_the_provider_under_a_generic_name(fake_transcribe):
    finished_job(post(audio=('Ward round with Dr X.webm', WEBM)))
    assert fake_transcribe == [{'kind': 'webm', 'name': 'audio.webm', 'bytes': WEBM}]


def test_slides_without_text_get_a_warning():
    state = finished_job(post(files=[('scanned.pdf', text_pdf(''))]))
    assert state['materials']['warnings'] == [
        'scanned.pdf has no text Lucy can read. It may contain pictures of slides.'
    ]


def test_transcription_failure_is_explained(monkeypatch):
    def broken(audio, kind):
        raise providers.ProviderError('The transcription service is busy. Please try again in a minute.')

    monkeypatch.setattr(providers, 'transcribe', broken)
    state = finished_job(post(audio=('rec.webm', WEBM)))
    assert state == {
        'status': 'failed',
        'step': 'Failed',
        'materials': None,
        'pack': None,
        'error': 'The transcription service is busy. Please try again in a minute.',
    }


def test_uploads_are_deleted_when_the_job_ends(monkeypatch):
    folders = []
    original = jobs.new_folder
    monkeypatch.setattr(jobs, 'new_folder', lambda: folders.append(original()) or folders[-1])

    finished_job(post(files=[('a.pdf', text_pdf('Hello'))], audio=('rec.webm', WEBM)))
    assert len(folders) == 1
    assert not folders[0].exists()


def test_uploads_are_deleted_when_the_job_fails(monkeypatch):
    folders = []
    original = jobs.new_folder
    monkeypatch.setattr(jobs, 'new_folder', lambda: folders.append(original()) or folders[-1])

    def broken(audio, kind):
        raise providers.ProviderError('busy')

    def crashing(audio, kind):
        raise RuntimeError('unexpected')

    for failure in (broken, crashing):
        monkeypatch.setattr(providers, 'transcribe', failure)
        assert finished_job(post(audio=('rec.webm', WEBM)))['status'] == 'failed'
    assert len(folders) == 2
    assert not any(folder.exists() for folder in folders)


def test_unexpected_errors_get_a_plain_message(monkeypatch):
    def crashing(audio, kind):
        raise RuntimeError('internal detail that must not reach the student')

    monkeypatch.setattr(providers, 'transcribe', crashing)
    assert finished_job(post(audio=('rec.webm', WEBM)))['error'] == 'Something went wrong while reading your material.'


def test_unreadable_pptx_gets_a_warning_not_a_failure():
    broken = office_zip('ppt/presentation.xml')
    state = finished_job(post('notes', [('broken.pptx', broken)]))
    assert state['status'] == 'done'
    assert state['materials']['warnings'] == ['broken.pptx could not be read.']


def test_pptx_tables_groups_and_line_breaks_are_read():
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = 'Severity', 'PEF'
    table.cell(1, 0).text, table.cell(1, 1).text = 'Acute severe', '33 to 50%'
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text_frame.text = 'Oxygen\vSalbutamol'
    buffer = io.BytesIO()
    deck.save(buffer)

    text = finished_job(post(files=[('table.pptx', buffer.getvalue())]))['materials']['slides'][0]['text']
    assert 'Severity | PEF' in text
    assert 'Acute severe | 33 to 50%' in text
    assert 'Oxygen\nSalbutamol' in text


def test_finished_results_are_forgotten_once_collected():
    job_id = post('notes').json()['job_id']
    assert client.get(f'/api/jobs/{job_id}').status_code == 200
    assert client.get(f'/api/jobs/{job_id}').status_code == 404


def test_leftover_uploads_are_swept_at_startup():
    leftover = jobs.new_folder()
    (leftover / 'audio.webm').write_bytes(WEBM)
    with TestClient(app):
        pass
    assert not leftover.exists()


def test_transcription_can_use_its_own_provider():
    from app.config import Settings

    shared = Settings(llm_base_url='https://llm.example/v1', llm_api_key='llm-key')
    assert shared.transcription_provider == ('https://llm.example/v1', 'llm-key')
    separate = Settings(
        llm_base_url='https://llm.example/v1',
        llm_api_key='llm-key',
        transcribe_base_url='https://speech.example/v1',
        transcribe_api_key='speech-key',
    )
    assert separate.transcription_provider == ('https://speech.example/v1', 'speech-key')


def test_unknown_job_is_404():
    assert client.get(f'/api/jobs/{"0" * 32}').status_code == 404


def test_malformed_job_id_is_refused():
    assert client.get('/api/jobs/..%2F..%2Fsecrets').status_code in (404, 422)


def test_missing_key_is_explained_without_calling_out(tmp_path, monkeypatch):
    monkeypatch.undo()  # use the real transcribe; the test settings have no key
    audio = tmp_path / 'audio.webm'
    audio.write_bytes(WEBM)
    with pytest.raises(providers.ProviderError, match='no AI key'):
        providers.transcribe(audio, 'webm')


def test_cors_allows_only_the_front_end():
    allowed = client.get('/api/health', headers={'Origin': 'http://localhost:5173'})
    other = client.get('/api/health', headers={'Origin': 'https://evil.example'})
    assert allowed.headers['access-control-allow-origin'] == 'http://localhost:5173'
    assert 'access-control-allow-origin' not in other.headers


def test_body_size_limit():
    tiny = FastAPI()

    @tiny.post('/echo')
    async def echo(body: dict) -> dict:
        return body

    tiny.add_middleware(BodySizeLimit, max_bytes=10)
    limited = TestClient(tiny)
    assert limited.post('/echo', json={'a': 1}).status_code == 200
    assert limited.post('/echo', json={'a': 'x' * 50}).status_code == 413
    streamed = limited.post('/echo', content=iter([b'{"a": "', b'x' * 50, b'"}']))
    assert streamed.status_code == 413
