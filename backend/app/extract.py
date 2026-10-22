"""Text from slides, and photos made ready for reading. Files come from strangers, so every reader has limits."""

import io
import zipfile
from collections.abc import Iterator
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

MAX_PAGES = 300
# Office files are zips; refuse any that would unpack to more than this (a "zip bomb").
MAX_UNPACKED_BYTES = 300 * 1024 * 1024
# A small file can claim to be a huge picture (an "image bomb"); the size is read before any pixels are.
MAX_PHOTO_PIXELS = 50_000_000
# Enough to read handwriting, and small enough for the AI service's limit on images.
PHOTO_LONG_SIDE = 2048


class UnreadableFile(Exception):
    pass


def pdf_text(path: Path) -> tuple[str, int]:
    """Returns the text of the first MAX_PAGES pages, and how many pages the file has."""
    try:
        reader = PdfReader(path)
        text = '\n\n'.join(page.extract_text() or '' for page in reader.pages[:MAX_PAGES])
        return text.strip(), len(reader.pages)
    except Exception as error:
        raise UnreadableFile from error


def photo_for_reading(path: Path) -> bytes:
    """A smaller JPEG copy, turned the right way up. Saving a new copy also drops what the camera
    stored inside the photo, such as the place and time it was taken."""
    try:
        with Image.open(path) as photo:
            if photo.width * photo.height > MAX_PHOTO_PIXELS:
                raise UnreadableFile
            upright = ImageOps.exif_transpose(photo)
            upright.thumbnail((PHOTO_LONG_SIDE, PHOTO_LONG_SIDE))
            out = io.BytesIO()
            upright.convert('RGB').save(out, 'JPEG', quality=85)
            return out.getvalue()
    except UnreadableFile:
        raise
    except Exception as error:
        raise UnreadableFile from error


def pptx_text(path: Path) -> tuple[str, int]:
    """Returns the text of the first MAX_PAGES slides (with speaker notes), and how many slides there are."""
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise UnreadableFile
        slides = list(Presentation(str(path)).slides)
    except UnreadableFile:
        raise
    except Exception as error:
        raise UnreadableFile from error

    parts = []
    for number, slide in enumerate(slides[:MAX_PAGES], start=1):
        lines = [line for shape in slide.shapes for line in _shape_text(shape)]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = _clean(slide.notes_slide.notes_text_frame.text)
            if notes:
                lines.append(f'Speaker notes: {notes}')
        if lines:
            parts.append(f'Slide {number}\n' + '\n'.join(lines))
    return '\n\n'.join(parts), len(slides)


def _shape_text(shape) -> Iterator[str]:
    # Clinical slides often keep their content in tables and grouped shapes, not plain text boxes.
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for inner in shape.shapes:
            yield from _shape_text(inner)
    elif getattr(shape, 'has_table', False) and shape.has_table:
        for row in shape.table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            if any(cells):
                yield ' | '.join(cells)
    elif shape.has_text_frame and _clean(shape.text_frame.text):
        yield _clean(shape.text_frame.text)


def _clean(text: str) -> str:
    # PowerPoint stores a soft line break (Shift+Enter) as a vertical tab.
    return text.replace('\v', '\n').strip()
